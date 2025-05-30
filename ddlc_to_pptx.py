import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import argparse
import os

class RenPyToPowerPoint:
    def __init__(self, json_file, output_pptx):
        self.json_file = json_file
        self.output_pptx = output_pptx
        self.presentation = Presentation()
        self.data = None
        self.slide_layouts = {
            'title': 0,
            'content': 1,
            'section_header': 2,
            'two_content': 3,
            'comparison': 4,
            'title_only': 5,
            'blank': 6
        }
        
        # Default styling
        self.title_font_size = Pt(36)
        self.subtitle_font_size = Pt(24)
        self.body_font_size = Pt(18)
        self.character_color = RGBColor(79, 129, 189)  # Blue
        self.narrator_color = RGBColor(128, 0, 128)    # Purple
        self.menu_color = RGBColor(34, 139, 34)        # Green
        self.bg_color = RGBColor(240, 240, 240)        # Light gray
        self.text_color = RGBColor(0, 0, 0)            # Black

    def load_data(self):
        """Load the JSON data from the Ren'Py parser"""
        with open(self.json_file, 'r', encoding='utf-8') as f:
            self.data = json.load(f)
            
    def create_title_slide(self):
        """Create the title slide for the presentation"""
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[self.slide_layouts['title']]
        )
        
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "Ren'Py Visual Novel"
        subtitle.text = "Scene Breakdown Presentation"
        
        # Style the title
        title.text_frame.paragraphs[0].font.size = self.title_font_size
        title.text_frame.paragraphs[0].font.bold = True
        
        # Style the subtitle
        subtitle.text_frame.paragraphs[0].font.size = self.subtitle_font_size

    def create_overview_slide(self):
        """Create an overview slide with project statistics"""
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[self.slide_layouts['content']]
        )
        
        title = slide.shapes.title
        title.text = "Project Overview"
        
        content = slide.placeholders[1]
        tf = content.text_frame
        
        # Add project statistics
        stats = [
            f"Total Scenes: {len(self.data['scenes'])}",
            f"Characters: {len(self.data['characters'])}",
            f"Images: {len(self.data['images'])}",
            f"Audio Files: {len(self.data['audio'])}"
        ]
        
        for stat in stats:
            p = tf.add_paragraph()
            p.text = stat
            p.font.size = self.body_font_size
            p.level = 0

    def create_character_slide(self):
        """Create a slide listing all characters"""
        if not self.data['characters']:
            return
            
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[self.slide_layouts['two_content']]
        )
        
        title = slide.shapes.title
        title.text = "Characters"
        
        left_content = slide.placeholders[1]
        right_content = slide.placeholders[2]
        
        left_tf = left_content.text_frame
        right_tf = right_content.text_frame
        
        # Split characters into two columns
        char_items = list(self.data['characters'].items())
        half = len(char_items) // 2
        
        for i, (var, name) in enumerate(char_items):
            tf = left_tf if i < half else right_tf
            p = tf.add_paragraph()
            p.text = f"{var}: {name}"
            p.font.size = self.body_font_size - Pt(2)
            p.level = 0

    def create_scene_slides(self):
        """Create slides for each scene"""
        for i, scene in enumerate(self.data['scenes']):
            self.create_single_scene_slide(scene, i+1)
            
    def create_single_scene_slide(self, scene, scene_num):
        """Create a slide for a single scene"""
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[self.slide_layouts['title_only']]
        )
        
        # Set slide title with scene number and label path
        title = slide.shapes.title
        title.text = f"Scene {scene_num}"
        
        # Add label path as subtitle if available
        if 'label_path' in scene and scene['label_path']:
            subtitle = slide.placeholders[1]
            subtitle.text = " → ".join(scene['label_path'])
            subtitle.text_frame.paragraphs[0].font.size = self.subtitle_font_size - Pt(4)
            subtitle.text_frame.paragraphs[0].font.italic = True
        
        # Create a text box for the content
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(5)
        text_box = slide.shapes.add_textbox(left, top, width, height)
        tf = text_box.text_frame
        
        # Add background information if available
        if scene.get('background'):
            p = tf.add_paragraph()
            p.text = f"Background: {scene['background']}"
            p.font.size = self.body_font_size - Pt(2)
            p.font.bold = True
            p.level = 0
        
        # Add sprites information if available
        if scene.get('sprites'):
            p = tf.add_paragraph()
            p.text = "Sprites:"
            p.font.size = self.body_font_size - Pt(2)
            p.font.bold = True
            p.level = 0
            
            for pos, sprite in scene['sprites'].items():
                p = tf.add_paragraph()
                p.text = f"  - {sprite['image']} at {pos}"
                p.font.size = self.body_font_size - Pt(2)
                p.level = 1
        
        # Add dialogue if available
        if scene.get('dialogue'):
            p = tf.add_paragraph()
            p.text = "Dialogue:"
            p.font.size = self.body_font_size - Pt(2)
            p.font.bold = True
            p.level = 0
            
            for line in scene['dialogue']:
                p = tf.add_paragraph()
                if line['speaker']:
                    p.text = f"{line['speaker']}: {line['text']}"
                    p.font.color = self.character_color
                else:
                    p.text = f"Narrator: {line['text']}"
                    p.font.color = self.narrator_color
                p.font.size = self.body_font_size - Pt(2)
                p.level = 1
        
        # Add audio if available
        if scene.get('audio'):
            p = tf.add_paragraph()
            p.text = "Audio:"
            p.font.size = self.body_font_size - Pt(2)
            p.font.bold = True
            p.level = 0
            
            for audio in scene['audio']:
                p = tf.add_paragraph()
                p.text = f"  - {audio}"
                p.font.size = self.body_font_size - Pt(2)
                p.level = 1
        
        # Add menu choices if available
        if scene.get('choices'):
            p = tf.add_paragraph()
            p.text = "Menu Choices:"
            p.font.size = self.body_font_size - Pt(2)
            p.font.bold = True
            p.font.color = self.menu_color
            p.level = 0
            
            for choice in scene['choices']:
                p = tf.add_paragraph()
                p.text = f"  - {choice['text']}"
                if choice.get('target'):
                    p.text += f" → {choice['target']}"
                p.font.size = self.body_font_size - Pt(2)
                p.font.color = self.menu_color
                p.level = 1

    def create_assets_slides(self):
        """Create slides listing all assets"""
        if self.data['images']:
            self.create_asset_slide("Images", self.data['images'])
        if self.data['audio']:
            self.create_asset_slide("Audio Files", self.data['audio'])
            
    def create_asset_slide(self, title, assets):
        """Create a slide for a specific type of asset"""
        # Split assets into chunks to avoid too much text on one slide
        chunk_size = 20
        for i in range(0, len(assets), chunk_size):
            slide = self.presentation.slides.add_slide(
                self.presentation.slide_layouts[self.slide_layouts['two_content']]
            )
            
            slide_title = slide.shapes.title
            slide_title.text = f"{title} (Part {i//chunk_size + 1})"
            
            left_content = slide.placeholders[1]
            right_content = slide.placeholders[2]
            
            left_tf = left_content.text_frame
            right_tf = right_content.text_frame
            
            chunk = assets[i:i+chunk_size]
            half = len(chunk) // 2
            
            for j, asset in enumerate(chunk):
                tf = left_tf if j < half else right_tf
                p = tf.add_paragraph()
                p.text = asset
                p.font.size = self.body_font_size - Pt(2)
                p.level = 0

    def generate_presentation(self):
        """Generate the complete PowerPoint presentation"""
        print("Loading JSON data...")
        self.load_data()
        
        print("Creating title slide...")
        self.create_title_slide()
        
        print("Creating overview slide...")
        self.create_overview_slide()
        
        print("Creating character slides...")
        self.create_character_slide()
        
        print(f"Creating {len(self.data['scenes'])} scene slides...")
        self.create_scene_slides()
        
        print("Creating asset slides...")
        self.create_assets_slides()
        
        print(f"Saving presentation to {self.output_pptx}...")
        self.presentation.save(self.output_pptx)
        print("Presentation created successfully!")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert Ren'Py JSON output to PowerPoint")
    parser.add_argument("json_file", help="Path to JSON file generated by RenPySceneParser")
    parser.add_argument("output_pptx", help="Output PowerPoint file path (.pptx)")
    args = parser.parse_args()

    if not os.path.exists(args.json_file):
        print(f"Error: JSON file not found at {args.json_file}")
        exit(1)

    converter = RenPyToPowerPoint(args.json_file, args.output_pptx)
    converter.generate_presentation()
